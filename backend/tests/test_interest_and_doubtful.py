"""Interest-brand aliases and illegible LED ranges."""

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from app import db
from app.config.interest_brands import resolve_interest_fragment
from app.pipeline.aggregate import FrameObservation
from app.pipeline.brands import match_brand_ids, prepare_brands
from app.pipeline.catalog import assignable_brands_from_meta
from app.pipeline.catalog_report import build_catalog_report
from app.pipeline.doubtful import aggregate_unmeasurable
from app.pipeline.led_timing import LedTiming
from app.pipeline.report import write_commercial_report
from app.schemas import BrandInput, BrandResult, DoubtfulSegment


LED = LedTiming(merge_gap_sec=8.0, on_confirm_sec=0.0, off_hold_sec=8.0)


def _obs(
    index: int,
    *,
    brands: frozenset[str] = frozenset(),
    ocr: str = "",
    ambiguous: frozenset[str] = frozenset(),
    skipped: bool = False,
    zone_id: str | None = "led",
    tipo_panel: str | None = "LED",
    half: str = "1T",
) -> FrameObservation:
    return FrameObservation(
        half=half,
        time_seconds=float(index),
        frame_idx=index,
        detected_brand_ids=brands,
        skipped=skipped,
        zone_id=zone_id,
        tipo_panel=tipo_panel,
        ocr_text=ocr,
        ambiguous_brand_ids=ambiguous,
    )


class InterestAliasTests(unittest.TestCase):
    def test_ocr_fragments_resolve_to_canonical_ids(self) -> None:
        self.assertEqual(resolve_interest_fragment("NETT OIUS").id, "nettplus")
        self.assertEqual(resolve_interest_fragment("ECUEBET").id, "ecuabet")
        self.assertEqual(resolve_interest_fragment("LIONS").id, "lions")
        self.assertEqual(resolve_interest_fragment("1 X BET").id, "1xbet")
        self.assertIsNone(resolve_interest_fragment("LIGAECUABET"))

    def test_prepared_brand_uses_config_aliases_not_a_hardcoded_list(self) -> None:
        brands = prepare_brands(
            [
                BrandInput(id="lions", name="LIONS"),
                BrandInput(id="1xbet", name="1xbet"),
                BrandInput(id="ecuabet", name="Ecuabet"),
            ]
        )
        self.assertEqual(
            match_brand_ids("LIONS SPORTS AND MEDIA LIONS SPORTS AND MEDIA", brands),
            {"lions"},
        )
        self.assertEqual(match_brand_ids("1 X BET 1 X BET", brands), {"1xbet"})
        self.assertEqual(match_brand_ids("ECUEBET ECUEBET", brands), {"ecuabet"})

    def test_fixed_board_matcher_is_not_required_for_alias_resolution(self) -> None:
        brands = prepare_brands([BrandInput(id="nettplus", name="NETTPLUS")])
        hits = match_brand_ids("NETT OIUS NETT OLUS", brands)
        self.assertEqual(hits, {"nettplus"})


class UnmeasurableRangeTests(unittest.TestCase):
    def test_illegible_stretch_is_one_wall_clock_range(self) -> None:
        rows = [_obs(index, ocr="XQZ QZX") for index in (0, 5)]
        ranges = aggregate_unmeasurable(rows, timing=LED)
        self.assertEqual(len(ranges), 1)
        self.assertEqual(ranges[0].duration_seconds, 6)
        self.assertEqual(ranges[0].reason, "illegible_ocr")
        self.assertFalse(ranges[0].reason_label.startswith("medible"))
        self.assertIn("no medible", ranges[0].reason_label)

    def test_measured_brand_interval_is_not_also_doubtful(self) -> None:
        rows = [
            _obs(0, brands=frozenset({"nett"})),
            _obs(1, brands=frozenset({"nett"})),
            _obs(3, ocr="????"),
            _obs(8, brands=frozenset({"nett"})),
        ]
        ranges = aggregate_unmeasurable(
            rows,
            timing=LED,
            covers=[("1T", 0.0, 9.0)],
        )
        self.assertEqual(ranges, [])

    def test_illegible_outside_a_brand_run_stays_for_review(self) -> None:
        rows = [
            _obs(0, brands=frozenset({"nett"})),
            _obs(1, brands=frozenset({"nett"})),
            _obs(10, ocr="GLPH"),
            _obs(12, ambiguous=frozenset({"ecuabet"})),
        ]
        ranges = aggregate_unmeasurable(
            rows,
            timing=LED,
            covers=[("1T", 0.0, 2.0)],
        )
        self.assertEqual(len(ranges), 1)
        self.assertEqual(ranges[0].start_seconds, 10.0)
        self.assertEqual(ranges[0].duration_seconds, 3)
        self.assertGreater(ranges[0].end_seconds, 2.0)

    def test_closeup_without_a_zone_is_not_a_review_range(self) -> None:
        rows = [_obs(index, skipped=True, zone_id=None) for index in range(5)]
        self.assertEqual(aggregate_unmeasurable(rows, timing=LED), [])

    def test_fixed_panel_is_ignored(self) -> None:
        rows = [_obs(index, ocr="NOISE", tipo_panel="FIXED_PRINT") for index in range(4)]
        self.assertEqual(aggregate_unmeasurable(rows, timing=LED), [])

    def test_low_quality_crop_on_a_located_zone_is_unmeasurable(self) -> None:
        rows = [_obs(0, skipped=True, zone_id="led-left"), _obs(1, skipped=True, zone_id="led-left")]
        ranges = aggregate_unmeasurable(rows, timing=LED)
        self.assertEqual(len(ranges), 1)
        self.assertEqual(ranges[0].reason, "low_led_quality")


class ReportClaimTests(unittest.TestCase):
    def test_excel_keeps_doubtful_seconds_out_of_the_led_total(self) -> None:
        brand = BrandResult(
            brand_id="nettplus",
            name="NETTPLUS",
            appearances=1,
            total_seconds=4,
            seconds=4,
            segments=[],
        )
        doubtful = DoubtfulSegment(
            half="1T",
            clock_start="00:10",
            clock_end="00:16",
            video_seconds_start=10.0,
            video_seconds_end=16.0,
            duration_seconds=6,
            reason="illegible_ocr",
            reason_label="no medible: OCR ilegible",
            ocr_text="XQZZ",
        )
        path = Path(tempfile.mkdtemp()) / "informe.xlsx"
        write_commercial_report(
            path,
            led_brands=[brand],
            analyzed_seconds=20,
            unmeasurable=[doubtful],
        )
        workbook = load_workbook(path)
        self.assertEqual(workbook["Resumen LED"].cell(row=2, column=3).value, 4)
        self.assertEqual(workbook["Resumen LED"].cell(row=2, column=8).value, "sí")
        reasons = [
            workbook["Dudosas"].cell(row=row, column=5).value
            for row in range(2, workbook["Dudosas"].max_row + 1)
        ]
        self.assertIn("no medible: OCR ilegible", reasons)
        workbook.close()

    def test_catalog_report_lists_interest_seconds_and_doubtful_ranges(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        db.set_db_path(Path(tmp.name) / "test.db")
        self.addCleanup(db.reset_db_path)
        db.init_db()
        db.upsert_brand("nett", "NETT plus")
        db.insert_job(
            job_id="job-doubt",
            stadium_id="ligaecuabet",
            profile_id=None,
            mode="single",
            duration_mode="full",
            status="completed",
            directory=tmp.name,
        )
        db.insert_job_frames(
            "job-doubt",
            [
                {
                    "half": "1T",
                    "frame_idx": 1,
                    "time_seconds": 0.0,
                    "crop_relpath": "a.jpg",
                    "machine_label": "positive",
                    "brand_id": "nett",
                    "ocr_text": "NETT PLUS",
                },
                {
                    "half": "1T",
                    "frame_idx": 2,
                    "time_seconds": 1.0,
                    "crop_relpath": "b.jpg",
                    "machine_label": "positive",
                    "brand_id": "nett",
                    "ocr_text": "NETT PLUS",
                },
                {
                    "half": "1T",
                    "frame_idx": 3,
                    "time_seconds": 20.0,
                    "crop_relpath": "c.jpg",
                    "machine_label": "attention",
                    "brand_id": None,
                    "ocr_text": "XQZZ QZX",
                },
                {
                    "half": "1T",
                    "frame_idx": 4,
                    "time_seconds": 21.0,
                    "crop_relpath": "d.jpg",
                    "machine_label": "attention",
                    "brand_id": None,
                    "ocr_text": "XQZZ",
                },
            ],
        )
        payload = build_catalog_report("job-doubt", timing=LED, sample_interval=1.0)
        assert payload is not None
        self.assertEqual(payload["summary"]["interest_seconds"], 2)
        self.assertEqual(payload["summary"]["total_seconds"], 2)
        self.assertTrue(payload["brands"][0]["interest"])
        self.assertEqual(payload["brands"][0]["canonical_name"], "NETTPLUS")
        self.assertEqual(payload["summary"]["doubtful_count"], 1)
        self.assertEqual(payload["summary"]["doubtful_seconds"], 2)
        segment = payload["doubtful_segments"][0]
        self.assertTrue(segment["doubtful"])
        self.assertFalse(segment["measurable"])
        self.assertEqual(segment["video_seconds_start"], 20.0)
        self.assertIn("no medible", segment["reason_label"])

    def test_assignable_list_includes_interest_brands_without_meta(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        rows = assignable_brands_from_meta(Path(tmp.name))
        ids = {row["brand_id"] for row in rows}
        self.assertIn("ecuabet", ids)
        self.assertIn("nettplus", ids)
        self.assertIn("lions", ids)
        self.assertIn("1xbet", ids)


if __name__ == "__main__":
    unittest.main()
