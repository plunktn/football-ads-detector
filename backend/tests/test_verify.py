import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import cv2
import numpy as np

from app.pipeline.aggregate import FrameObservation
from app.pipeline.playlist import (
    PlaylistSlot,
    libertad_orense_fixture_sheets,
    parse_playlist,
    write_lions_workbook,
)
from app.pipeline.report import write_commercial_report
from app.pipeline.run import run_analysis
from app.pipeline.stages import ShotKind
from app.pipeline.verify import (
    HIT_TOLERANCE_SEC,
    hit_rate,
    scheduled_video_sec,
    slot_video_window,
    to_compliance_row,
    verify_slot_from_observations,
)
from app.schemas import (
    BrandInput,
    BrandResult,
    Kickoff,
    SegmentResult,
    VerificationStatus,
)


class VerifySlotTests(unittest.TestCase):
    def test_maps_clock_to_video_using_kickoff(self) -> None:
        slot = PlaylistSlot("1T", "NETPLUS", 15, 30, 15)
        kickoff = Kickoff(
            first_half_video_seconds=10,
            second_half_video_seconds=100,
        )
        self.assertEqual(slot_video_window(slot, kickoff), (25.0, 40.0))
        second = PlaylistSlot("2T", "LIONS", 30, 45, 15)
        self.assertEqual(slot_video_window(second, kickoff), (130.0, 145.0))

    def test_real_match_offset_mapping(self) -> None:
        slot = PlaylistSlot("1T", "NETT.plus", 15, 30, 15)
        kickoff = Kickoff(
            first_half_video_seconds=282,
            second_half_video_seconds=3521,
        )
        self.assertEqual(scheduled_video_sec(slot, kickoff), 297.0)
        second = PlaylistSlot("2T", "LIONS", 30, 45, 15)
        self.assertEqual(scheduled_video_sec(second, kickoff), 3551.0)

    def test_past_eof_when_scheduled_beyond_duration(self) -> None:
        slot = PlaylistSlot("1T", "NETPLUS", 15, 30, 15)
        kickoff = Kickoff(first_half_video_seconds=282)
        result = verify_slot_from_observations(
            slot,
            [],
            kickoff,
            "netplus",
            video_duration_sec=200.0,
        )
        self.assertEqual(result.status, VerificationStatus.PAST_EOF)
        self.assertFalse(result.hit)

    def test_hit_from_observations_inside_window(self) -> None:
        slot = PlaylistSlot("1T", "NETPLUS", 15, 30, 15)
        kickoff = Kickoff(first_half_video_seconds=0)
        observations = [
            FrameObservation(
                half="1T",
                time_seconds=20.0,
                frame_idx=600,
                detected_brand_ids=frozenset({"netplus"}),
                skipped=False,
                tipo_panel="LED_DYNAMIC",
                shot=ShotKind.WIDE_LED,
            )
        ]
        result = verify_slot_from_observations(
            slot, observations, kickoff, "netplus"
        )
        self.assertEqual(result.status, VerificationStatus.HIT)
        self.assertTrue(result.hit)
        self.assertEqual(result.observed_video_sec, 20.0)
        self.assertEqual(result.delta_sec, 5.0)

    def test_miss_on_usable_led_without_brand(self) -> None:
        slot = PlaylistSlot("1T", "NETPLUS", 15, 30, 15)
        observations = [
            FrameObservation(
                half="1T",
                time_seconds=20.0,
                frame_idx=600,
                detected_brand_ids=frozenset(),
                skipped=False,
                tipo_panel="LED_DYNAMIC",
                shot=ShotKind.WIDE_LED,
                zone_id="lateral_main",
            )
        ]
        result = verify_slot_from_observations(
            slot, observations, Kickoff(), "netplus"
        )
        self.assertEqual(result.status, VerificationStatus.MISS)
        self.assertFalse(result.hit)

    def test_no_evidence_when_all_closeup_or_graphic(self) -> None:
        slot = PlaylistSlot("1T", "NETPLUS", 15, 30, 15)
        observations = [
            FrameObservation(
                half="1T",
                time_seconds=18.0,
                frame_idx=540,
                detected_brand_ids=frozenset(),
                skipped=True,
                shot=ShotKind.CLOSEUP,
            ),
            FrameObservation(
                half="1T",
                time_seconds=22.0,
                frame_idx=660,
                detected_brand_ids=frozenset(),
                skipped=True,
                shot=ShotKind.GRAPHIC_BUMPER,
            ),
        ]
        result = verify_slot_from_observations(
            slot, observations, Kickoff(), "netplus"
        )
        self.assertEqual(result.status, VerificationStatus.NO_EVIDENCE)
        self.assertFalse(result.hit)

    def test_offset_when_brand_outside_tolerance(self) -> None:
        slot = PlaylistSlot("1T", "NETPLUS", 15, 30, 15)
        # Window is 15–30; brand at 50 → Δ=35 > 20.
        observations = [
            FrameObservation(
                half="1T",
                time_seconds=20.0,
                frame_idx=600,
                detected_brand_ids=frozenset(),
                skipped=False,
                shot=ShotKind.WIDE_LED,
            ),
            FrameObservation(
                half="1T",
                time_seconds=50.0,
                frame_idx=1500,
                detected_brand_ids=frozenset({"netplus"}),
                skipped=False,
                shot=ShotKind.WIDE_LED,
                tipo_panel="LED_DYNAMIC",
            ),
        ]
        result = verify_slot_from_observations(
            slot, observations, Kickoff(), "netplus"
        )
        self.assertEqual(result.status, VerificationStatus.OFFSET)
        self.assertGreater(abs(result.delta_sec or 0), HIT_TOLERANCE_SEC)

    def test_ambiguous_weak_ocr(self) -> None:
        slot = PlaylistSlot("1T", "NETPLUS", 15, 30, 15)
        observations = [
            FrameObservation(
                half="1T",
                time_seconds=20.0,
                frame_idx=600,
                detected_brand_ids=frozenset(),
                skipped=False,
                shot=ShotKind.WIDE_LED,
                ambiguous_brand_ids=frozenset({"netplus"}),
                zone_id="lateral_main",
            )
        ]
        result = verify_slot_from_observations(
            slot, observations, Kickoff(), "netplus"
        )
        self.assertEqual(result.status, VerificationStatus.AMBIGUOUS)

    def test_hit_derived_on_compliance_row(self) -> None:
        slot = PlaylistSlot("1T", "NETPLUS", 15, 30, 15)
        result = verify_slot_from_observations(
            slot,
            [
                FrameObservation(
                    half="1T",
                    time_seconds=20.0,
                    frame_idx=1,
                    detected_brand_ids=frozenset({"netplus"}),
                    skipped=False,
                    shot=ShotKind.WIDE_LED,
                )
            ],
            Kickoff(),
            "netplus",
        )
        row = to_compliance_row(result)
        self.assertEqual(row.status, VerificationStatus.HIT)
        self.assertTrue(row.hit)

    def test_hit_rate_excludes_doubtful(self) -> None:
        rows = [
            to_compliance_row(
                verify_slot_from_observations(
                    PlaylistSlot("1T", "A", 0, 15, 15),
                    [
                        FrameObservation(
                            half="1T",
                            time_seconds=5.0,
                            frame_idx=1,
                            detected_brand_ids=frozenset({"a"}),
                            skipped=False,
                            shot=ShotKind.WIDE_LED,
                        )
                    ],
                    Kickoff(),
                    "a",
                )
            ),
            to_compliance_row(
                verify_slot_from_observations(
                    PlaylistSlot("1T", "B", 20, 35, 15),
                    [
                        FrameObservation(
                            half="1T",
                            time_seconds=25.0,
                            frame_idx=2,
                            detected_brand_ids=frozenset(),
                            skipped=False,
                            shot=ShotKind.WIDE_LED,
                        )
                    ],
                    Kickoff(),
                    "b",
                )
            ),
            to_compliance_row(
                verify_slot_from_observations(
                    PlaylistSlot("1T", "C", 40, 55, 15),
                    [
                        FrameObservation(
                            half="1T",
                            time_seconds=45.0,
                            frame_idx=3,
                            detected_brand_ids=frozenset(),
                            skipped=True,
                            shot=ShotKind.CLOSEUP,
                        )
                    ],
                    Kickoff(),
                    "c",
                )
            ),
        ]
        # 1 HIT + 1 MISS + 1 NO_EVIDENCE → 0.5
        self.assertEqual(hit_rate(rows), 0.5)


class PlaylistVerifyPipelineTests(unittest.TestCase):
    def _video(self, duration_seconds: int) -> Path:
        handle = tempfile.NamedTemporaryFile(suffix=".mp4", delete=False)
        handle.close()
        path = Path(handle.name)
        writer = cv2.VideoWriter(
            str(path),
            cv2.VideoWriter_fourcc(*"mp4v"),
            10,
            (64, 64),
        )
        for _ in range(duration_seconds * 10):
            writer.write(np.zeros((64, 64, 3), dtype=np.uint8))
        writer.release()
        self.addCleanup(path.unlink, missing_ok=True)
        return path

    def test_synthetic_video_hits_over_90_percent(self) -> None:
        from app.domain.zones import DEFAULT_PANEL_ZONES
        from app.pipeline.ocr import OcrHit
        from app.pipeline.roi import RoiResult

        video = self._video(40)
        tmp = Path(tempfile.mkdtemp())
        playlist = tmp / "PLAYLIST_LIBERTADvORENSE.xlsx"
        write_lions_workbook(
            playlist,
            {
                "PREVIA": [("LIONS", "0.00", 15)],
                "PRIMER TIEMPO": [
                    ("NETPLUS", "0.02", 5),
                    ("ECUABET", "0.08", 5),
                ],
                "ENTRETIEMPO": [],
                "SEGUNDO TIEMPO": [("NETPLUS", "0.02", 5)],
                "POST": [],
            },
        )
        slots = parse_playlist(playlist)
        crop = np.zeros((24, 80, 3), dtype=np.uint8)
        roi = RoiResult(False, None, crop, None, y0=10, y1=34)
        zone = DEFAULT_PANEL_ZONES[0]

        def fake_hits(_crop):
            return [
                OcrHit(text="NETPLUS", x_center=0.2),
                OcrHit(text="NETPLUS", x_center=0.8),
            ]

        with patch("app.pipeline.run.locate_zones", return_value=[(zone, roi)]), \
             patch("app.pipeline.run.read_led_hits", side_effect=fake_hits), \
             patch(
                 "app.pipeline.run.match_brand_ids",
                 side_effect=lambda *args, **kwargs: {"netplus", "ecuabet"},
             ), \
             patch("app.pipeline.run.match_fixed_brand_ids", return_value=set()), \
             patch(
                 "app.pipeline.run.classify_shot",
                 return_value=ShotKind.WIDE_LED,
             ):
            result = run_analysis(
                [video],
                mode="single",
                duration_mode="full",
                kickoff=Kickoff(
                    first_half_video_seconds=0,
                    second_half_video_seconds=20,
                ),
                brands=[
                    BrandInput(id="netplus", name="NETPLUS", aliases=["NETT.plus"]),
                    BrandInput(id="ecuabet", name="ECUABET"),
                ],
                debug_dir=tmp / "debug",
                analysis_mode="playlist_verify",
                playlist_slots=slots,
            )
        self.assertGreater(result.hit_rate or 0.0, 0.90)
        self.assertTrue(all(row.hit for row in result.compliance))
        self.assertTrue(
            all(row.status == VerificationStatus.HIT for row in result.compliance)
        )


class ReportTests(unittest.TestCase):
    def test_workbook_has_audit_sheets(self) -> None:
        tmp = Path(tempfile.mkdtemp()) / "informe.xlsx"
        brand = BrandResult(
            brand_id="netplus",
            name="NETPLUS",
            appearances=2,
            total_seconds=30,
            minutes=0,
            seconds=30,
            panel_kind="LED",
            count_1t=1,
            count_2t=1,
        )
        extra_brand = BrandResult(
            brand_id="mystery",
            name="MYSTERYCO",
            appearances=1,
            total_seconds=15,
            minutes=0,
            seconds=15,
            panel_kind="LED",
            count_1t=1,
            count_2t=0,
            segments=[
                SegmentResult(
                    half="1T",
                    clock_start="00:10",
                    clock_end="00:25",
                    video_seconds_start=10.0,
                    video_seconds_end=25.0,
                    start_frame=1,
                    end_frame=15,
                    duration_seconds=15,
                )
            ],
        )
        slots = parse_playlist(
            write_lions_workbook(
                tmp.with_name("playlist.xlsx"),
                libertad_orense_fixture_sheets(),
            )
        )
        hit_row = to_compliance_row(
            verify_slot_from_observations(
                next(s for s in slots if s.period == "1T"),
                [
                    FrameObservation(
                        half="1T",
                        time_seconds=20,
                        frame_idx=1,
                        detected_brand_ids=frozenset({"netplus"}),
                        skipped=False,
                        shot=ShotKind.WIDE_LED,
                    )
                ],
                Kickoff(),
                "netplus",
            )
        )
        doubtful = to_compliance_row(
            verify_slot_from_observations(
                PlaylistSlot("1T", "CAMERO", 150, 165, 15),
                [
                    FrameObservation(
                        half="1T",
                        time_seconds=155,
                        frame_idx=2,
                        detected_brand_ids=frozenset(),
                        skipped=True,
                        shot=ShotKind.CLOSEUP,
                    )
                ],
                Kickoff(),
                "camara decomercio",
            )
        )
        path = write_commercial_report(
            tmp,
            led_brands=[brand, extra_brand],
            fixed_brands=[],
            slots=slots,
            compliance=[hit_row, doubtful],
            hit_rate=1.0,
            analyzed_seconds=40,
        )
        from openpyxl import load_workbook

        wb = load_workbook(path)
        self.assertEqual(
            wb.sheetnames,
            [
                "Resumen",
                "Salidas",
                "Cumplimiento",
                "Dudosas",
                "Extras",
            ],
        )
        self.assertEqual(wb["Resumen"][1][0].value, "Cliente")
        self.assertEqual(wb["Resumen"][2][0].value, "NETPLUS")
        self.assertEqual(wb["Cumplimiento"][1][5].value, "Status")
        self.assertEqual(wb["Dudosas"][2][3].value, "NO_EVIDENCE")
        self.assertEqual(wb["Extras"][2][0].value, "MYSTERYCO")


if __name__ == "__main__":
    unittest.main()
