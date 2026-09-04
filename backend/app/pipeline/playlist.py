"""Parse Lions multi-sheet Excel playlists into timed brand slots."""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime, time
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook


DEFAULT_DURATION_SEC = 15.0

_HEADER_CLIENT = {"cliente", "clientes", "marca", "marcas", "client"}
_HEADER_MINUTE = {"minuto", "minutos", "min", "mmss", "hora"}
_HEADER_DURATION = {"duracion", "duración", "dur", "segundos", "seg", "duration"}

_PERIOD_ALIASES = {
    "PREVIA": "PREVIA",
    "PRE": "PREVIA",
    "PREVIAS": "PREVIA",
    "PRIMER TIEMPO": "1T",
    "PRIMER": "1T",
    "1T": "1T",
    "PT": "1T",
    "1ER TIEMPO": "1T",
    "ENTRETIEMPO": "ENTRETIEMPO",
    "ENTRE TIEMPO": "ENTRETIEMPO",
    "HT": "ENTRETIEMPO",
    "MEDIO TIEMPO": "ENTRETIEMPO",
    "SEGUNDO TIEMPO": "2T",
    "SEGUNDO": "2T",
    "2T": "2T",
    "ST": "2T",
    "2DO TIEMPO": "2T",
    "POST": "POST",
    "POSTPARTIDO": "POST",
    "POST PARTIDO": "POST",
}

_VERIFY_PERIODS = frozenset({"1T", "2T"})


@dataclass(frozen=True)
class PlaylistSlot:
    period: str
    brand: str
    start_sec: float
    end_sec: float
    duration_sec: float
    sheet: str = ""
    row: int = 0


def _strip_accents(value: str) -> str:
    return (
        unicodedata.normalize("NFKD", value)
        .encode("ascii", "ignore")
        .decode()
    )


def _norm_header(value: Any) -> str:
    if value is None:
        return ""
    text = _strip_accents(str(value)).strip().lower()
    text = re.sub(r"[^a-z0-9]+", "", text)
    return text


def _norm_sheet(value: str) -> str:
    text = _strip_accents(value).upper()
    text = re.sub(r"\s+", " ", text).strip()
    return text


def period_from_sheet(sheet_name: str) -> str:
    key = _norm_sheet(sheet_name)
    if key in _PERIOD_ALIASES:
        return _PERIOD_ALIASES[key]
    for alias, period in _PERIOD_ALIASES.items():
        if alias in key:
            return period
    return key or "UNKNOWN"


def parse_minuto(value: Any) -> float:
    """Parse Lions minute cells.

    ``0.15`` / ``"0.15"`` / ``"M.SS"`` / ``"MM.SS"`` → minutes:seconds.
    ``0:15`` and Excel time objects also work.
    """
    if value is None or value == "":
        raise ValueError("minuto vacío")
    if isinstance(value, time):
        return float(value.hour * 3600 + value.minute * 60 + value.second)
    if isinstance(value, datetime):
        return float(value.hour * 3600 + value.minute * 60 + value.second)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        if number < 0:
            raise ValueError(f"minuto inválido: {value!r}")
        # Whole numbers are clock minutes (15 → 15:00).
        if float(number).is_integer():
            return float(int(number) * 60)
        minutes = int(number)
        seconds = int(round((abs(number) - abs(minutes)) * 100))
        if seconds >= 60:
            raise ValueError(f"minuto inválido: {value!r}")
        return float(minutes * 60 + seconds)

    text = str(value).strip()
    if not text:
        raise ValueError("minuto vacío")
    if ":" in text:
        parts = text.split(":")
        if len(parts) < 2:
            raise ValueError(f"minuto inválido: {value!r}")
        minutes = int(parts[0])
        seconds = int(float(parts[1]))
        if seconds >= 60:
            raise ValueError(f"minuto inválido: {value!r}")
        return float(minutes * 60 + seconds)
    if "." in text:
        left, right = text.split(".", 1)
        minutes = int(left or "0")
        seconds = int(right or "0")
        if seconds >= 60:
            raise ValueError(f"minuto inválido: {value!r}")
        return float(minutes * 60 + seconds)
    return float(int(text) * 60)


def parse_duration(value: Any, default: float = DEFAULT_DURATION_SEC) -> float:
    if value is None or value == "":
        return default
    if isinstance(value, time):
        return float(value.hour * 3600 + value.minute * 60 + value.second)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        if number <= 0:
            return default
        return number
    text = str(value).strip().replace(",", ".")
    if not text:
        return default
    try:
        number = float(text)
    except ValueError as exc:
        raise ValueError(f"duración inválida: {value!r}") from exc
    return number if number > 0 else default


def _find_header_map(rows: Iterable[tuple[Any, ...]]) -> tuple[int, dict[str, int]] | None:
    for index, row in enumerate(rows):
        mapping: dict[str, int] = {}
        for col, cell in enumerate(row):
            key = _norm_header(cell)
            if key in _HEADER_CLIENT and "client" not in mapping:
                mapping["client"] = col
            elif key in _HEADER_MINUTE and "minute" not in mapping:
                mapping["minute"] = col
            elif key in _HEADER_DURATION and "duration" not in mapping:
                mapping["duration"] = col
        if "client" in mapping and "minute" in mapping:
            return index, mapping
    return None


def parse_playlist(path: str | Path) -> list[PlaylistSlot]:
    """Parse every sheet of a Lions xlsx. Never only ``wb.active``."""
    workbook = load_workbook(filename=str(path), data_only=True, read_only=True)
    slots: list[PlaylistSlot] = []
    try:
        for sheet in workbook.worksheets:
            period = period_from_sheet(sheet.title)
            rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
            header = _find_header_map(rows[:12])
            if header is None:
                continue
            header_index, mapping = header
            client_col = mapping["client"]
            minute_col = mapping["minute"]
            duration_col = mapping.get("duration")
            for offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
                if not row:
                    continue
                client_raw = row[client_col] if client_col < len(row) else None
                if client_raw is None or str(client_raw).strip() == "":
                    continue
                minute_raw = row[minute_col] if minute_col < len(row) else None
                try:
                    start_sec = parse_minuto(minute_raw)
                except (TypeError, ValueError):
                    continue
                duration_raw = (
                    row[duration_col]
                    if duration_col is not None and duration_col < len(row)
                    else None
                )
                duration_sec = parse_duration(duration_raw)
                slots.append(
                    PlaylistSlot(
                        period=period,
                        brand=str(client_raw).strip(),
                        start_sec=start_sec,
                        end_sec=start_sec + duration_sec,
                        duration_sec=duration_sec,
                        sheet=sheet.title,
                        row=offset,
                    )
                )
    finally:
        workbook.close()
    return slots


def slots_for_verify(slots: Iterable[PlaylistSlot]) -> list[PlaylistSlot]:
    return [slot for slot in slots if slot.period in _VERIFY_PERIODS]


def unique_brands(slots: Iterable[PlaylistSlot]) -> list[str]:
    seen: dict[str, str] = {}
    for slot in slots:
        key = slot.brand.strip()
        if key and key.lower() not in seen:
            seen[key.lower()] = key
    return list(seen.values())


def write_lions_workbook(
    path: str | Path,
    sheets: dict[str, list[tuple[str, Any, Any]]],
) -> Path:
    """Write a multi-sheet Lions-style xlsx. Values: (cliente, minuto, duracion)."""
    destination = Path(path)
    workbook = Workbook()
    first = True
    for sheet_name, rows in sheets.items():
        if first:
            sheet = workbook.active
            assert sheet is not None
            sheet.title = sheet_name
            first = False
        else:
            sheet = workbook.create_sheet(sheet_name)
        sheet.append(["CLIENTE", "MINUTO", "DURACIÓN"])
        for client, minute, duration in rows:
            sheet.append([client, minute, duration])
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()
    return destination


def libertad_orense_fixture_sheets() -> dict[str, list[tuple[str, Any, Any]]]:
    """Minimal copy of PLAYLIST_LIBERTADvORENSE (Lions columns, all five periods)."""
    return {
        "PREVIA": [
            ("LIONS SPORTS AND MEDIA", "0.00", 15),
        ],
        "PRIMER TIEMPO": [
            ("NETT.plus", "0.15", 15),
            ("ECUABET", "1.00", 15),
            ("CAMERO", "2.30", 15),
            ("PILSENER", "5.00", 15),
        ],
        "ENTRETIEMPO": [
            ("HOTEL VICTORIA", "0.00", 15),
        ],
        "SEGUNDO TIEMPO": [
            ("NETPLUS", "0.00", 15),
            ("LIONS", "0.30", 15),
            ("COBLERED", "1.15", 15),
        ],
        "POST": [
            ("UTPL", "0.00", 15),
        ],
    }
